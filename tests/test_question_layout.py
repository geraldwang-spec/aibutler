import json
import re
import tempfile
import unittest
from pathlib import Path

from personal_ai.question_importer import extract_by_rules, extract_questions
from personal_ai.question_layout import extract_layout, source_fragments
from personal_ai.llm_provider import LLMError


class LayoutModel:
    """Return source locations, never generated question wording."""
    max_input_bytes = 16000
    calls = 0
    field = 'content'

    def complete_json(self, system, user):
        self.calls += 1
        data = json.loads(user.split('\n回傳 ')[0])
        runs = []
        for key, text in data['fragments'].items():
            index = int(key)
            option = re.match(r'^([AB])\.', text)
            if option:
                self.field = option[1]
            elif text.startswith('答案：'):
                self.field = 'answer'
            elif text.startswith('解析：'):
                self.field = 'explanation'
            if runs and runs[-1]['field'] == self.field:
                runs[-1]['e'] = index
            else:
                runs.append(dict(s=index,e=index,field=self.field,new=self.calls==1 and index==0,type='單選'))
        return {'parts': runs}


class QuestionLayoutTests(unittest.TestCase):
    def test_unknown_passage_scope_is_resolved_after_later_questions(self):
        class Deferred:
            max_input_bytes=16000
            field='content'
            scope_calls=0
            def complete_json(self,system,user):
                data=json.JSONDecoder().raw_decode(user)[0]
                if 'questions' in data:
                    self.scope_calls+=1
                    assert '後面的子題' in system
                    assert [q['number'] for q in data['questions']]==[2,3]
                    return {'question_numbers':[2]}
                parts=[]
                for index,text in data['fragments'].items():
                    new=False
                    if text.startswith('Reading passage'):
                        self.field='passage'
                        new=True
                    elif re.match(r'^[123]\.',text): self.field='content'; new=True
                    elif re.match(r'^[AB]\.',text): self.field=text[0]
                    elif text.startswith('答案：'): self.field='answer'
                    parts.append(dict(s=int(index),e=int(index),field=self.field,new=new,type='單選',question_from=None,question_to=None))
                return {'parts':parts}
        text='1. Previous?\nA. one\nB. two\n答案：A\nReading passage about a farm.\n2. What does the farm provide?\nA. food\nB. traffic\n答案：A\n3. Independent arithmetic question?\nA. one\nB. two\n答案：B'
        model=Deferred()
        rows=extract_layout(text,model,'測試')
        self.assertEqual(len(rows),3)
        self.assertEqual(model.scope_calls,1)
        self.assertIn('Reading passage',rows[1]['content'])
        self.assertNotIn('Reading passage',rows[0]['content'])
        self.assertNotIn('Reading passage',rows[2]['content'])
        self.assertEqual(rows[1]['_passage_range'],[2,2])

    def test_unconfirmed_passage_is_preserved_without_aborting_other_questions(self):
        from personal_ai.question_layout import _resolve_passage_scopes
        class Uncertain:
            max_input_bytes=16000
            def complete_json(self,s,u): return {'question_numbers':[]}
        items=[dict(_question_no=1,content='獨立題'),dict(_question_no=2,content='閱讀子題')]
        passages=[dict(position=1,question_from=None,question_to=None,text='完整原文')]
        _resolve_passage_scopes(passages,items,Uncertain())
        self.assertNotIn('_layout_needs_review',items[0])
        self.assertIn('_layout_needs_review',items[1])
        self.assertEqual(passages[0]['text'],'完整原文')

    def test_groq_json_400_shrinks_layout_and_preserves_source(self):
        class JsonFailure(LayoutModel):
            sizes=[]
            def complete_json(self,system,user):
                data=json.loads(user.split('\n回傳 ')[0])
                self.sizes.append(len(data['fragments']))
                if len(data['fragments'])>2:
                    raise LLMError('JSON rejected',status_code=400,error_code='json_validate_failed')
                assert '鍵與字串用雙引號' in user
                return super().complete_json(system,user)
        model=JsonFailure()
        rows=extract_layout('1. Choose.\nA. apple\nB. book\n答案：B\n解析：理由',model,'測試')
        self.assertGreater(model.sizes[0],2)
        self.assertEqual(rows[0]['option_A'],'apple')
        self.assertEqual(rows[0]['option_B'],'book')
        self.assertEqual(rows[0]['answer_key'],'B')

    def test_groq_json_400_persistent_failure_is_bounded(self):
        class Broken:
            max_input_bytes=16000
            calls=0
            def complete_json(self,system,user):
                self.calls+=1
                raise LLMError('bad JSON',status_code=400,error_code='json_validate_failed')
        model=Broken()
        with self.assertRaisesRegex(LLMError,'有限重試'):
            extract_layout('1. Question',model,'測試')
        self.assertEqual(model.calls,2)

    def test_unrelated_http_400_is_not_retried_as_json_failure(self):
        class WrongRequest:
            max_input_bytes=16000
            calls=0
            def complete_json(self,system,user):
                self.calls+=1
                raise LLMError('bad model',status_code=400,error_code='model_not_found')
        model=WrongRequest()
        with self.assertRaises(LLMError): extract_layout('1. Q\nA. x\nB. y',model,'測試')
        self.assertEqual(model.calls,1)

    def test_semantic_layout_attaches_shared_passage_without_special_heading(self):
        class ReadingModel:
            max_input_bytes=16000
            field='content'
            calls=0
            def complete_json(self,system,user):
                self.calls+=1
                assert '最後選項不是吸收' in system
                assert 'passage' in system and '克漏字' in system
                data=json.loads(user.split('\n回傳 ')[0])
                parts=[]
                for index,text in data['fragments'].items():
                    new=False
                    if text.startswith('材料'):
                        self.field='passage'
                    elif re.match(r'^(?:1|12|13)\.',text):
                        self.field='content'
                        new=True
                    elif re.match(r'^[AB]\.',text):
                        self.field=text[0]
                    elif text.startswith('答案：'):
                        self.field='answer'
                    part=dict(s=int(index),e=int(index),field=self.field,new=new,type='單選')
                    if self.field=='passage': part.update(question_from=12,question_to=13)
                    parts.append(part)
                return {'parts':parts}
        passage='材料：'+('The city farm provides food and community activities. '*85)+'\n'
        text='1. Prior question\nA. prior A\nB. prior B\n答案：A\n'+passage+'12. What is provided?\nA. food\nB. traffic\n答案：A\n13. What else is provided?\nA. traffic\nB. community activities\n答案：B'
        model=ReadingModel()
        rows=extract_layout(text,model,'閱讀')
        self.assertGreater(model.calls,1)
        self.assertEqual([x['_question_no'] for x in rows],[1,12,13])
        self.assertNotIn('city farm',rows[0]['content'])
        self.assertEqual(rows[0]['option_B'],'prior B')
        for row in rows[1:]:
            self.assertIn(passage.strip(),row['content'])
            self.assertTrue(row['_shared_passage'])
            self.assertEqual(row['_passage_range'],[12,13])

    def test_auto_uses_valid_local_structure_but_explicit_ai_checks_layout(self):
        from unittest.mock import patch
        sections=[{'text':'1. Question\n(A)one(B)two\n答案：A'}]
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'fixture.txt'
            path.write_text(sections[0]['text'],encoding='utf-8')
            parsed=extract_by_rules(sections,'測試')
            for mode in ('auto','llm'):
                with patch('personal_ai.question_importer.extract_with_llm',return_value=parsed) as layout:
                    result,_=extract_questions(path,'測試',{},mode)
                    self.assertEqual(result,parsed)
                    self.assertEqual(layout.call_count,0 if mode=='auto' else 1)

    def test_pdf_reading_early_number_includes_preceding_source_page(self):
        import fitz
        from unittest.mock import patch
        from personal_ai.question_answering import infer_missing_answers
        class Vision:
            enabled=True
            provider='test'
            model='vision'
            images=[]
            def complete_json_with_images(self,system,user,images):
                self.images=images
                assert '比較題必須讀到所有比較材料' in system
                assert '附圖來源頁碼（依順序）：1,2' in user
                return {'answers':[dict(number=3,status='answered',answer='B',explanation='原文支持',context='')]}
        vision=Vision()
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'reading.pdf'
            with fitz.open() as doc:
                doc.new_page().insert_text((40,60),'Shared reading passage on preceding page.')
                doc.new_page().insert_text((40,60),'3. Read and choose.\n(A) one\n(B) two')
                doc.save(path)
            item=dict(_question_no=3,content='Read and choose.',q_type='單選',answer_key='',option_A='one',option_B='two',_shared_passage=True,_passage_range=[3,4])
            with patch('personal_ai.question_answering.get_parser_llm',return_value=vision),patch('personal_ai.question_answering.get_vision_llm',return_value=vision):
                result=infer_missing_answers([item],[],path,{})
            self.assertEqual(result[0]['answer_key'],'B')
            self.assertEqual(len(vision.images),2)

    def test_inline_next_passage_does_not_become_option_d(self):
        text='''25. 兩首詩的共同點？
(A)描寫食物 (B)用典 (C)顏色對比 (D)以食材入詩，強調美好事物稍縱即逝 請閱讀以下資料，並回答26～29題：【甲】城市農場的文章。
【乙】海上牧場的文章。
26. 甲文的主旨？
(A)城市農場 (B)海上牧場 (C)交通 (D)能源
27. 乙文的主旨？
(A)城市農場 (B)海上牧場 (C)交通 (D)能源
30. 另一題？
(A)甲 (B)乙 (C)丙 (D)丁'''
        rows=extract_by_rules([{'text':text}],'國文',allow_missing_answers=True)
        self.assertEqual([x['_question_no'] for x in rows],[25,26,27,30])
        self.assertEqual(rows[0]['option_D'],'以食材入詩，強調美好事物稍縱即逝')
        self.assertNotIn('農場',rows[0]['content'])
        for row in rows[1:3]:
            self.assertIn('【甲】城市農場的文章',row['content'])
            self.assertIn('【乙】海上牧場的文章',row['content'])
            self.assertNotIn('文章',row['option_D'])
        self.assertNotIn('農場',rows[3]['content'])

    def test_passage_heading_on_own_line_and_fullwidth_range(self):
        text='1. 題幹\n(A)甲(B)乙(C)丙(D)閱讀是好習慣\n請閱讀下文，回答第２～３題：共享文章\n2. 第一問\n(A)甲(B)乙\n3. 第二問\n(A)丙(B)丁'
        rows=extract_by_rules([{'text':text}],'測試',allow_missing_answers=True)
        self.assertEqual(rows[0]['option_D'],'閱讀是好習慣')
        self.assertIn('共享文章',rows[1]['content'])
        self.assertIn('共享文章',rows[2]['content'])

    def test_layout_shrinks_on_truncation_without_losing_source(self):
        class Limited(LayoutModel):
            def complete_json(self,system,user):
                payload=json.loads(user.split('\n回傳 ')[0])
                if len(payload['fragments'])>2:
                    raise LLMError('truncated',error_code='output_truncated')
                return super().complete_json(system,user)
        source='1. Choose.\nA. apple\nB. book\n答案：B\n解析：理由'
        rows=extract_layout(source,Limited(),'測試')
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['option_B'],'book')
        self.assertEqual(rows[0]['answer_key'],'B')

    def test_layout_rejects_incomplete_runs_before_mutation(self):
        class Repair(LayoutModel):
            failed=False
            def complete_json(self,system,user):
                if not self.failed:
                    self.failed=True
                    return {'parts':[dict(s=0,e=0,field='content',new=True,type='單選')]}
                return super().complete_json(system,user)
        rows=extract_layout('1. Choose.\nA. apple\nB. book\n答案：B',Repair(),'測試')
        self.assertEqual(rows[0]['content'],'Choose.')
        self.assertEqual(len(rows),1)

    def test_answer_aliases_and_numeric_choices(self):
        from personal_ai.question_layout import split_fields,answer_text
        options,stem,answer,_=split_fields('Choose.\n（一）apple\n（二）book\n解答：二')
        self.assertEqual(options,{'A':'apple','B':'book'})
        self.assertEqual(answer,'B')
        self.assertEqual(answer_text('Correct Answer: ○'),'是')
        self.assertEqual(answer_text('答案為：×'),'否')

    def test_extra_options_are_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'book.txt'
            path.write_text('1. Choose\nA. one\nB. two\nC. three\nD. four\nE. five\n答案：A',encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'E／F'):
                extract_questions(path,'測試',{},'rules')

    def test_inference_shrinks_truncated_batch(self):
        from unittest.mock import patch
        from personal_ai.question_answering import infer_missing_answers
        class Solver:
            enabled=True
            provider='test'
            model='test'
            sizes=[]
            def complete_json(self,s,u):
                batch=json.loads(u.split('\n回傳 ')[0])
                self.sizes.append(len(batch))
                if len(batch)>1: raise LLMError('truncated',error_code='output_truncated')
                return {'answers':[dict(number=batch[0]['number'],status='answered',answer='B',explanation='理由',context='')]}
        model=Solver()
        items=[dict(_question_no=i,content='乙',q_type='單選',answer_key='',option_A='a',option_B='b') for i in (1,2)]
        with patch('personal_ai.question_answering.get_parser_llm',return_value=model):
            rows=infer_missing_answers(items,[],Path('book.txt'),{})
        self.assertEqual(model.sizes,[2,1,1])
        self.assertEqual([r['_question_no'] for r in rows],[1,2])

    def test_inferred_answers_do_not_replace_official_answers(self):
        from unittest.mock import patch
        from personal_ai.question_answering import infer_missing_answers
        class Solver:
            enabled=True
            provider='test'
            model='test'
            def complete_json(self,s,u):
                return {'answers':[dict(number=2,status='answered',answer='B',explanation='理由',context='')]}
        official=dict(_question_no=1,content='甲',q_type='單選',answer_key='A',option_A='a',option_B='b')
        missing=dict(_question_no=2,content='乙',q_type='單選',answer_key='',option_A='a',option_B='b')
        with patch('personal_ai.question_answering.get_parser_llm',return_value=Solver()):
            result=infer_missing_answers([official,missing],[],Path('book.txt'),{})
        self.assertIs(result[0],official)
        self.assertEqual(result[1]['answer_key'],'B')

    def test_rate_limit_retry_honors_provider_cooldown(self):
        from unittest.mock import patch,Mock
        from personal_ai.question_answering import _limited_request
        call=Mock(side_effect=[LLMError('rate limit',status_code=429,retry_after='2'),{'ok':True}])
        with patch('personal_ai.question_answering.time.sleep') as sleep:
            self.assertEqual(_limited_request(call),{'ok':True})
        self.assertEqual(call.call_count,2)
        self.assertEqual(sleep.call_count,3)

    def test_ocr_step_numbers_do_not_become_question_numbers(self):
        sections=[{'text':'（ ）20. Choose.\n(A)tea\n(B)water',
                   'image_context':[{'range':(20,21),'text':'1. Boil water.\n2. Add tea.'}]}]
        rows=extract_by_rules(sections,'測試',allow_missing_answers=True)
        self.assertEqual([x['_question_no'] for x in rows],[20])
        self.assertIn('1. Boil water.',rows[0]['content'])

    def test_question_book_without_answers_reports_missing_answer(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'book.txt'
            path.write_text('（   ）1. Choose a word.\n(A)apple\n(B)book',encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'1 個題號.*標準答案'):
                extract_questions(path,'測試',{},'rules')

    def test_missing_answer_inference_is_labelled(self):
        from unittest.mock import patch
        from personal_ai.question_answering import infer_missing_answers,AI_PREFIX
        class Solver:
            enabled=True
            provider='test'
            model='test-solver'
            def complete_json(self,system,user):
                return {'answers':[dict(number=1,status='answered',answer='B',explanation='依題意判斷。',context='')]}
        row=dict(_question_no=1,content='題目',q_type='單選',answer_key='',option_A='甲',option_B='乙')
        with patch('personal_ai.question_answering.get_parser_llm',return_value=Solver()):
            result=infer_missing_answers([row],[],Path('test.txt'),{})
        self.assertEqual(result[0]['answer_key'],'B')
        self.assertEqual(result[0]['_answer_source'],'ai_inferred')
        self.assertTrue(result[0]['explanation'].startswith(AI_PREFIX))
        self.assertEqual(row['answer_key'],'')

    def test_missing_answer_solver_refuses_unreadable_material(self):
        from unittest.mock import patch
        from personal_ai.question_answering import infer_missing_answers
        class Solver:
            enabled=True
            provider='test'
            model='test-solver'
            def complete_json(self,system,user):
                return {'answers':[dict(number=1,status='insufficient',answer='',explanation='圖片不清楚',context='')]}
        with patch('personal_ai.question_answering.get_parser_llm',return_value=Solver()):
            with self.assertRaisesRegex(LLMError,'足夠可讀材料'):
                infer_missing_answers([dict(_question_no=1,content='看圖',q_type='單選',option_A='甲',option_B='乙')],[],Path('test.txt'),{})

    def test_background_and_question_are_not_split_without_answer(self):
        class SplitModel:
            max_input_bytes=16000
            def complete_json(self,system,user):
                return {'parts':[
                    dict(s=0,e=0,field='content',new=True,type='單選'),
                    dict(s=1,e=1,field='content',new=True,type='單選'),
                    dict(s=2,e=2,field='A',new=False),
                    dict(s=3,e=3,field='B',new=False),
                    dict(s=4,e=4,field='answer',new=False)]}
        rows=extract_layout('背景文章。\n請選出正確敘述。\n(A)甲\n(B)乙\n答案：A',SplitModel(),'測試')
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['content'],'背景文章。\n請選出正確敘述。')

    def test_numeric_options_are_not_question_numbers(self):
        from personal_ai.question_importer import _question_start
        self.assertIsNone(_question_start('(A)20 mL'))
        self.assertIsNone(_question_start('(B) 30 mL'))
        self.assertEqual(_question_start('( B ) 12. 題目'), 12)
        rows=self.rules('1. 容量多少？\n(A)20 mL\n(B)30 mL\n答案：B')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['option_A'], '20 mL')

    def test_docx_order_and_broken_pdf_text_detection(self):
        from docx import Document
        from personal_ai.parsers import parse_file, needs_ocr
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'order.docx'
            doc=Document()
            doc.add_paragraph('第一題')
            doc.add_table(rows=1,cols=1).cell(0,0).text='第一題選項'
            doc.add_paragraph('第二題')
            doc.save(path)
            self.assertEqual([s['text'] for s in parse_file(path)], ['第一題','第一題選項','第二題'])
        self.assertTrue(needs_ocr('(cid:1)(cid:2)(cid:3)'+'文字'*30))
        self.assertFalse(needs_ocr('正常的題目文字'*20))

    def rules(self, text):
        return extract_by_rules([{'text':text}], '測試')

    def test_fullwidth_parenthesized_and_multiline_explanation(self):
        rows = self.rules('題目１：下列何者正確？\n（Ａ）選項甲\n（Ｂ）選項乙\n【答案】Ｂ\n【解析】第一行\n第二行')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['answer_key'], 'B')
        self.assertEqual(rows[0]['option_A'], '選項甲')
        self.assertEqual(rows[0]['explanation'], '第一行\n第二行')

    def test_unnumbered_blocks(self):
        rows = self.rules('第一題內容\nA. 甲\nB. 乙\n答案：A\n\n第二題內容\nA. 丙\nB. 丁\n答案：B')
        self.assertEqual([r['answer_key'] for r in rows], ['A','B'])

    def test_passage_is_kept_for_each_group_question(self):
        rows = self.rules('(40-41)\nA long time ago, there was a king.\n40. 選擇第一題\n(A) yes\n(B) no\n答案：B\n41. 選擇第二題\n(A) yes\n(B) no\n答案：A')
        self.assertEqual(len(rows), 2)
        self.assertTrue(all('there was a king' in r['content'] for r in rows))

    def test_long_source_is_copied_without_model_reproduction(self):
        content = '閱讀長文。' * 1400 + 'TAIL_SENTINEL'
        source = '1. '+content+'\nA. 甲\nB. 乙\n答案：B\n解析：原文解析'
        model = LayoutModel()
        rows = extract_layout(source, model, '長文')
        self.assertGreater(model.calls, 1)
        self.assertEqual(rows[0]['content'], content)
        self.assertEqual(rows[0]['answer_key'], 'B')
        self.assertEqual(''.join(source_fragments(source)), source)

    def test_incomplete_model_coverage_is_rejected(self):
        class Broken:
            def complete_json(self, *args):
                return {'parts':[dict(s=0,e=0,field='content',new=True)]}
        with self.assertRaises(LLMError):
            extract_layout('題目\nA. 甲\nB. 乙\n答案：B',Broken(),'測試')

    def test_partial_rules_are_not_silently_imported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'questions.txt'
            path.write_text('1. 題目\nA. 甲\nB. 乙\n答案：A\n2. 題目沒有答案\nA. 甲\nB. 乙',encoding='utf-8')
            rows, _ = extract_questions(path,'測試',{},mode='rules')
            self.assertEqual(rows, [])

    def test_long_question_validation_has_no_application_size_guard(self):
        from personal_ai.question_validation import validate
        row = dict(q_type='單選',content='長'*12000,answer_key='A',explanation='解析')
        validated, pairs = validate(row, {'A':'甲','B':'乙'})
        self.assertEqual(validated['content'], row['content'])
        long_row,_=validate(dict(row,content='😀'*16000), {'A':'甲','B':'乙'})
        self.assertEqual(long_row['content'],'😀'*16000)

    def test_long_classification_sees_tail_and_preserves_original(self):
        from personal_ai.classification_context import prepare
        class Summarizer:
            enabled=True
            provider='test'
            fragments=[]
            def complete_json(self,system,user):
                self.fragments.append(user)
                return {'hint':'原文的知識點'}
        model=Summarizer()
        item=dict(content='長'*12000+'TAIL_SENTINEL',answer_key='A',option_A='甲',option_B='乙')
        context=prepare(item,model)
        self.assertIn('TAIL_SENTINEL',''.join(model.fragments))
        self.assertTrue(item['content'].endswith('TAIL_SENTINEL'))
        self.assertTrue(context['_long_classification_context'])

    def test_long_question_reaches_import_preview_storage(self):
        import io
        from flask import g
        from werkzeug.datastructures import FileStorage
        from app import create_app
        from storage import db
        from smartlife import process_import_file
        with tempfile.TemporaryDirectory() as directory:
            app = create_app({'TESTING':True,'DATABASE':str(Path(directory)/'test.db'),
                              'SECRET_KEY':'layout-test','DB_READ_ONLY':False})
            app.instance_path = directory
            with app.test_request_context('/imports',method='POST'):
                db().execute("INSERT INTO users(id,username,email,password_hash) VALUES(1,'tester','tester@example.com','unused')")
                db().execute("INSERT INTO subjects(id,subject_name,created_by) VALUES(1,'測試',1)")
                db().commit()
                g.user = {'id':1}
                content='長文'*6000+'TAIL_SENTINEL'
                upload=FileStorage(stream=io.BytesIO(('1. '+content+'\n(A) 甲\n(B) 乙\n答案：B').encode('utf-8')),filename='long.txt')
                batch=process_import_file(app,{'subject_id':'1','import_strategy':'question_bank','parse_mode':'rules'},upload)
                row=db().execute('SELECT parsed_json FROM import_items WHERE import_id=?',(batch,)).fetchone()
                self.assertEqual(json.loads(row[0])['content'], content)


if __name__ == '__main__':
    unittest.main()
