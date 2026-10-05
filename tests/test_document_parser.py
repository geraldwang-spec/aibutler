"""共用文件轉換（modules/document_parser.py）的測試：測試檔案在記憶體中產生，不需要外部檔案。"""
import io
import unittest

from modules.document_parser import DocumentParseError, DocumentParser


def sample_docx():
    from docx import Document
    document = Document()
    document.add_heading('教練建議', level=1)
    document.add_paragraph('這份文件說明下個月的訓練重點。')
    document.add_heading('背部', level=2)
    document.add_paragraph('每週安排划船和下拉，平衡推的動作。', style='List Bullet')
    table = document.add_table(rows=3, cols=3)
    for r, values in enumerate([('動作', '組數', '次數'), ('槓鈴划船', '4', '10'), ('滑輪下拉', '3', '12')]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    document.add_paragraph('表格之後的說明。')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_xlsx():
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '週課表'
    for row in [('星期', '動作', '組數', '次數'), ('一', '槓鈴划船', 4, 10), ('三', '槓鈴深蹲', 5, 5), (None, None, None, None)]:
        sheet.append(row)
    empty = workbook.create_sheet('空白')
    empty.append([None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def sample_pdf(pages, text_on_page=True):
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    document = pymupdf.open()
    for number in range(1, pages + 1):
        page = document.new_page()
        if text_on_page:
            page.insert_text((72, 60), '考試智伴教練講義', fontname='china-t', fontsize=10)          # 每頁重複的頁首
            page.insert_text((72, 120), f'第{number}頁：背部訓練要安排划船與下拉，每週十組以上。', fontname='china-t', fontsize=12)
            page.insert_text((72, 800), f'- {number} -', fontsize=9)
    data = document.tobytes()
    document.close()
    return data


class DocumentParserTests(unittest.TestCase):
    def test_text_and_big5(self):
        self.assertEqual(DocumentParser.parse('a.md', '# 標題\n\n內容'.encode('utf-8'))[0]['text'], '# 標題\n\n內容')
        self.assertEqual(DocumentParser.parse('a.txt', '背部訓練'.encode('big5'))[0]['text'], '背部訓練')
        self.assertEqual(DocumentParser.parse('a.txt', '\ufeff有 BOM'.encode('utf-8'))[0]['text'], '有 BOM')

    def test_csv_and_excel_rows_become_sentences(self):
        csv_text = '動作,組數,次數\n槓鈴划船,4,10\n,,\n'.encode('utf-8')
        self.assertEqual(DocumentParser.parse('a.csv', csv_text)[0]['text'], '動作：槓鈴划船；組數：4；次數：10')
        sections = DocumentParser.parse('plan.xlsx', sample_xlsx())
        self.assertEqual([s['locator'] for s in sections], ['工作表「週課表」'])            # 空白工作表略過
        self.assertIn('星期：一；動作：槓鈴划船；組數：4；次數：10', sections[0]['text'])

    def test_table_without_header(self):
        self.assertEqual(DocumentParser.table_to_text([['1', '2'], ['3', '4']]), '1｜2\n3｜4')

    def test_docx_keeps_order_headings_lists_tables(self):
        text = DocumentParser.parse('coach.docx', sample_docx())[0]['text']
        self.assertTrue(text.startswith('# 教練建議'))
        self.assertIn('## 背部', text)
        self.assertIn('- 每週安排划船和下拉', text)
        self.assertLess(text.index('動作：槓鈴划船；組數：4；次數：10'), text.index('表格之後的說明'))

    def test_pdf_pages_and_repeated_header(self):
        sections = DocumentParser.parse('notes.pdf', sample_pdf(3))
        self.assertEqual([s['locator'] for s in sections], ['第 1 頁', '第 2 頁', '第 3 頁'])
        self.assertIn('第2頁：背部訓練要安排划船與下拉', sections[1]['text'])
        self.assertNotIn('考試智伴教練講義', sections[1]['text'])                          # 每頁重複的頁首被拿掉

    def test_errors_are_friendly(self):
        with self.assertRaises(DocumentParseError) as ctx:
            DocumentParser.parse('scan.pdf', sample_pdf(2, text_on_page=False))
        self.assertIn('掃描檔', str(ctx.exception))
        for name, data, words in (('a.exe', b'x', '支援'), ('a.txt', b'', '空'), ('a.docx', b'not zip', '無法開啟'),
                                  ('a.xlsx', b'not zip', '無法開啟'), ('a.pdf', b'%PDF-broken', '無法開啟'),
                                  ('a.txt', b'\xff\xfe\xfa\xfb' * 3, '編碼')):
            with self.assertRaises(DocumentParseError) as ctx:
                DocumentParser.parse(name, data)
            self.assertIn(words, str(ctx.exception), name)


if __name__ == '__main__':
    unittest.main()
