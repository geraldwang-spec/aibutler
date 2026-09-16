import csv
import io
import json
import math
import os
import secrets
import sqlite3
import zipfile
from xml.etree.ElementTree import ParseError
from datetime import date, datetime, timedelta
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, abort, flash, g, jsonify, redirect, render_template, request, session, url_for, Response
from werkzeug.security import generate_password_hash

from auth import auth, login_required
from records import CATALOG, PROFILE_FIELDS, options, ownership
from storage import db, init_storage


def create_app(test_config=None):
    root = Path(__file__).parent
    load_dotenv(root / '.env')
    app = Flask(__name__)
    app.config.update(DATABASE=str(root/'instance'/'smartlife.db'), SECRET_KEY=os.getenv('SECRET_KEY',''),
                      MAX_CONTENT_LENGTH=5*1024*1024, SESSION_COOKIE_HTTPONLY=True,
                      SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE')=='true',
                      PERMANENT_SESSION_LIFETIME=timedelta(hours=8))
    for key in ('SMTP_HOST','SMTP_PORT','SMTP_USERNAME','SMTP_PASSWORD','SMTP_FROM_EMAIL'):
        app.config[key] = os.getenv(key,'')
    app.config['SMTP_FROM_NAME'] = os.getenv('SMTP_FROM_NAME','考試智伴')
    app.config['SMTP_SECURITY'] = os.getenv('SMTP_SECURITY','starttls')
    if test_config:
        app.config.update(test_config)
    if not app.config['SECRET_KEY']:
        secret_file = root/'instance'/'session.key'
        secret_file.parent.mkdir(exist_ok=True)
        try:
            with secret_file.open('x', encoding='utf-8') as file:
                file.write(secrets.token_hex(32))
        except FileExistsError:
            pass
        app.config['SECRET_KEY'] = secret_file.read_text(encoding='utf-8').strip()
    app.config['DUMMY_PASSWORD_HASH'] = generate_password_hash(secrets.token_hex(16))
    init_storage(app)
    app.register_blueprint(auth)

    @app.before_request
    def security():
        g.user = None
        if session.get('user_id'):
            user = db().execute('SELECT * FROM users WHERE id=?',(session['user_id'],)).fetchone()
            if user and user['is_email_verified'] and session.get('password_changed_at') == user['password_changed_at']:
                g.user = user
            else:
                session.clear()
        if 'csrf_token' not in session:
            session['csrf_token'] = secrets.token_hex(32)
        if request.method == 'POST' and not secrets.compare_digest(request.form.get('csrf_token',''), session['csrf_token']):
            if request.path == '/send-verification':
                return jsonify(ok=False,message='頁面已過期，請重新整理後再試。'),400
            abort(400, '頁面已過期，請重新整理後再試。')

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'SAMEORIGIN'
        response.headers['Referrer-Policy'] = 'same-origin'
        response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self'; img-src 'self' data:; form-action 'self'; frame-ancestors 'self'"
        if request.endpoint != 'static':
            response.headers['Cache-Control'] = 'no-store'
        return response

    @app.context_processor
    def common():
        return dict(catalog=CATALOG, csrf_token=session.get('csrf_token',''), today=date.today().isoformat())

    @app.errorhandler(400)
    @app.errorhandler(404)
    @app.errorhandler(413)
    def error_page(error):
        message = '檔案不可超過 5 MB。' if error.code==413 else ('找不到此資料，或您沒有存取權限。' if error.code==404 else str(error.description))
        return render_template('error.html',title='無法完成操作',message=message), error.code

    @app.get('/')
    def index():
        return redirect(url_for('dashboard' if g.user else 'auth.login'))

    @app.get('/dashboard')
    @login_required
    def dashboard():
        refresh_stats()
        counts = [db().execute('SELECT count(*) FROM subjects WHERE created_by=?',(g.user['id'],)).fetchone()[0],
                  db().execute('SELECT count(*) FROM quiz_sessions WHERE user_id=? AND finished_at IS NOT NULL',(g.user['id'],)).fetchone()[0],
                  db().execute("SELECT count(*) FROM wrong_answers WHERE user_id=? AND status='待複習'",(g.user['id'],)).fetchone()[0]]
        plans = db().execute('SELECT * FROM study_plans WHERE user_id=? ORDER BY plan_date DESC LIMIT 10',(g.user['id'],)).fetchall()
        daily = db().execute('SELECT * FROM daily_summary WHERE user_id=? ORDER BY summary_date DESC LIMIT 14',(g.user['id'],)).fetchall()
        return render_template('dashboard_live.html',title='學習與健康總覽',counts=counts,plans=plans,daily=daily)

    @app.route('/profile',methods=['GET','POST'])
    @login_required
    def profile():
        row = db().execute('SELECT * FROM user_profiles WHERE user_id=?',(g.user['id'],)).fetchone()
        values = dict(row) if row else {}
        error = None
        if request.method=='POST':
            values = request.form
            try:
                data = parse_fields(PROFILE_FIELDS)
                if data['birth_date'] > date.today().isoformat():
                    raise ValueError('生日不可晚於今天。')
                columns = ','.join(data)
                updates = ','.join(k+'=excluded.'+k for k in data)
                db().execute(f'INSERT INTO user_profiles (user_id,{columns},onboarded_at) VALUES (?,{",".join("?" for _ in data)},CURRENT_TIMESTAMP) ON CONFLICT(user_id) DO UPDATE SET {updates},onboarded_at=CURRENT_TIMESTAMP',(g.user['id'],*data.values()))
                initial_weight = request.form.get('initial_weight','').strip()
                if initial_weight:
                    weight = float(initial_weight)
                    if not math.isfinite(weight) or not 1<=weight<=600:
                        raise ValueError('體重須介於 1–600 kg。')
                    db().execute('INSERT INTO body_metrics (user_id,record_date,weight_kg) VALUES (?,?,?) ON CONFLICT(user_id,record_date) DO UPDATE SET weight_kg=excluded.weight_kg',(g.user['id'],date.today().isoformat(),weight))
                db().commit()
                flash('個人資料已儲存。','success')
                return redirect(url_for('dashboard'))
            except (ValueError,sqlite3.IntegrityError) as exc:
                db().rollback()
                error = str(exc) if isinstance(exc,ValueError) else '資料格式錯誤，請重新檢查。'
        return render_template('profile.html',title='個人資料與目標',fields=PROFILE_FIELDS,values=values,error=error,choices={})

    @app.route('/records/<table>',methods=['GET','POST'])
    @login_required
    def records(table):
        if table not in CATALOG:
            abort(404)
        title, icon, fields = CATALOG[table]
        error = None
        values = {}
        edit_id = request.args.get('edit',type=int)
        if edit_id:
            values = dict(owned(table,edit_id))
        if request.method=='POST':
            values = request.form
            try:
                record_id = request.form.get('id',type=int)
                if record_id:
                    owned(table,record_id)
                data = parse_fields(fields)
                if table == 'questions':
                    # Freeze questions used in exams; changing answers would corrupt history.
                    if record_id and db().execute('SELECT 1 FROM quiz_answers WHERE question_id=?',(record_id,)).fetchone():
                        raise ValueError('此題已有作答紀錄，請新增另一題以保留歷史版本。')
                    data, question_options = validate_question(data,request.form)
                if table == 'workouts':
                    start,end = datetime.fromisoformat(data['started_at']),datetime.fromisoformat(data['ended_at'])
                    if end<=start or start.date().isoformat()!=data['workout_date'] or (end-start).total_seconds()>86400:
                        raise ValueError('訓練開始日期須符合訓練日，結束須晚於開始且時長不超過 24 小時。')
                    data['duration_min'] = int((end-start).total_seconds()/60)
                if table=='template_items':
                    template=owned('workout_templates',data['template_id'])
                    if data['day_index']>template['days_per_week']:
                        raise ValueError('訓練日不可超過課表的一週天數。')
                if table=='workout_templates':
                    if record_id and db().execute('SELECT 1 FROM template_items WHERE template_id=? AND day_index>?',(record_id,data['days_per_week'])).fetchone():
                        raise ValueError('請先調整超過新天數的課表項目。')
                    if data['is_active']=='1':
                        db().execute('UPDATE workout_templates SET is_active=0 WHERE user_id=?',(g.user['id'],))
                if table=='chat_messages':
                    data['role']='user'
                if record_id:
                    db().execute(f'UPDATE {table} SET '+','.join(k+'=?' for k in data)+' WHERE id=?',(*data.values(),record_id))
                else:
                    if table in ('subjects','exercises'):
                        data['created_by']=g.user['id']
                    elif table not in ('chapters','questions','workout_sets','template_items','chat_messages'):
                        data['user_id']=g.user['id']
                    record_id = insert(table,data)
                if table=='questions':
                    db().execute('DELETE FROM question_options WHERE question_id=?',(record_id,))
                    for label,text in question_options:
                        insert('question_options',dict(question_id=record_id,option_label=label,option_text=text,order_no=ord(label)))
                refresh_stats(commit=False)
                db().commit()
                flash('資料已儲存。','success')
                return redirect(url_for('records',table=table))
            except (ValueError,sqlite3.IntegrityError) as exc:
                db().rollback()
                error = str(exc) if isinstance(exc,ValueError) else '資料重複或關聯不正確；同一天的體重請編輯原紀錄。'
        rows = db().execute(f'SELECT * FROM {table} WHERE {ownership(table)} ORDER BY id DESC',(g.user['id'],)).fetchall()
        choices = {field.name:options(db(),field.ref,g.user['id']) for field in fields if field.ref}
        option_values = {}
        if table=='questions' and edit_id:
            option_values = {'option_'+r['option_label']:r['option_text'] for r in db().execute('SELECT * FROM question_options WHERE question_id=?',(edit_id,))}
        return render_template('records.html',title=title,table=table,fields=fields,rows=rows,choices=choices,values=values,error=error,edit_id=edit_id,option_values=option_values)

    @app.post('/records/<table>/<int:record_id>/delete')
    @login_required
    def delete_record(table,record_id):
        if table not in CATALOG:
            abort(404)
        owned(table,record_id)
        try:
            if table=='questions':
                db().execute('DELETE FROM question_options WHERE question_id=?',(record_id,))
            db().execute(f'DELETE FROM {table} WHERE id=?',(record_id,))
            refresh_stats(commit=False)
            db().commit()
            flash('資料已刪除。','success')
        except sqlite3.IntegrityError:
            db().rollback()
            flash('此資料仍被其他紀錄使用，請先處理相關資料；已有作答的題目會保留。','error')
        return redirect(url_for('records',table=table))

    @app.get('/workspace/<page>')
    @login_required
    def workspace(page):
        mapping={'overview':'dashboard','mock':'quiz_start','results':'results','analysis':'analysis','sources':'imports','wellness':'dashboard'}
        if page=='review':
            return redirect(url_for('records',table='summaries'))
        if page not in mapping:
            abort(404)
        return redirect(url_for(mapping[page]))

    for old, endpoint in [('data-visualization','analysis'),('maps','quiz_start'),('manage-users','profile'),('preferences','profile')]:
        app.add_url_rule('/'+old,old,login_required(lambda target=endpoint: redirect(url_for(target))))

    register_learning(app)
    return app


def insert(table,data):
    return db().execute(f'INSERT INTO {table} ({",".join(data)}) VALUES ({",".join("?" for _ in data)})',tuple(data.values())).lastrowid


def owned(table,record_id):
    row=db().execute(f'SELECT * FROM {table} WHERE id=? AND {ownership(table)}',(record_id,g.user['id'])).fetchone()
    if not row:
        abort(404)
    return row


def parse_fields(fields):
    data={}
    for field in fields:
        value=request.form.get(field.name,'').strip()
        if not value:
            if field.required:
                raise ValueError('請填寫：'+field.label)
            data[field.name]=None
            continue
        if field.kind in ('number','decimal','ref'):
            try:
                value=float(value) if field.kind=='decimal' else int(value)
            except ValueError:
                raise ValueError(field.label+' 必須是有效數字。')
            if field.ref:
                owned(field.ref,value)
            elif not math.isfinite(value) or not field.minimum<=value<=field.maximum:
                raise ValueError(f'{field.label} 須介於 {field.minimum}–{field.maximum}。')
        elif field.kind=='select':
            if value not in field.choices:
                raise ValueError(field.label+' 選項無效。')
        elif field.kind in ('date','datetime-local'):
            try:
                parsed=date.fromisoformat(value) if field.kind=='date' else datetime.fromisoformat(value)
                if field.kind=='datetime-local' and parsed.tzinfo is not None:
                    raise ValueError()
                value=parsed.isoformat()
            except ValueError:
                raise ValueError(field.label+' 日期或時間無效。')
        elif len(value)>field.limit:
            raise ValueError(f'{field.label} 不可超過 {field.limit} 字。')
        data[field.name]=value
    return data


def validate_question(data,form):
    opts=[(c,str(form.get('option_'+c,'') or '').strip()) for c in 'ABCD']
    if any(len(text)>2000 for _,text in opts):
        raise ValueError('每個選項不可超過 2000 字。')
    opts=[(c,text) for c,text in opts if text]
    answer=data['answer_key'].strip()
    if data['q_type'] in ('單選','多選'):
        answer=','.join(sorted(set(answer.upper().replace('，',',').replace(' ','').split(','))))
        labels=answer.split(',')
        if len(opts)<2 or not set(labels)<=set(c for c,_ in opts) or (data['q_type']=='單選' and len(labels)!=1):
            raise ValueError('選擇題至少填兩個選項，正確答案須使用已填的選項代號（多選用 A,C）。')
    elif data['q_type']=='是非':
        if answer not in ('是','否'):
            raise ValueError('是非題答案請填「是」或「否」。')
        opts=[]
    else:
        opts=[]
    data['answer_key']=answer
    return data,opts


def refresh_stats(commit=True):
    uid=g.user['id']
    db().execute("UPDATE study_plans SET status='missed' WHERE user_id=? AND status='planned' AND plan_date<?",(uid,date.today().isoformat()))
    db().execute('UPDATE workouts SET total_volume=COALESCE((SELECT sum(weight_kg*reps) FROM workout_sets WHERE workout_id=workouts.id),0) WHERE user_id=?',(uid,))
    daily={}
    for r in db().execute('SELECT substr(a.answered_at,1,10) day,count(*) n,sum(a.is_correct) c FROM quiz_answers a JOIN quiz_sessions s ON s.id=a.session_id WHERE s.user_id=? AND s.finished_at IS NOT NULL GROUP BY day',(uid,)):
        daily[r['day']]={'answered_count':r['n'],'accuracy':round(r['c']/r['n']*100,2)}
    for r in db().execute('SELECT workout_date,sum(total_volume) volume FROM workouts WHERE user_id=? GROUP BY workout_date',(uid,)):
        daily.setdefault(r['workout_date'],{})['workout_volume']=r['volume']
    db().execute('DELETE FROM daily_summary WHERE user_id=?',(uid,))
    for day,data in daily.items():
        groups=[r[0] for r in db().execute('SELECT DISTINCT e.muscle_group FROM workout_sets s JOIN workouts w ON w.id=s.workout_id JOIN exercises e ON e.id=s.exercise_id WHERE w.user_id=? AND w.workout_date=?',(uid,day))]
        insert('daily_summary',dict(user_id=uid,summary_date=day,trained_groups=json.dumps(groups,ensure_ascii=False),**data))
    if commit:
        db().commit()


def register_learning(app):
    @app.route('/quiz',methods=['GET','POST'])
    @login_required
    def quiz_start():
        error=None
        if request.method=='POST':
            try:
                subject_id=int(request.form.get('subject_id','0'))
                owned('subjects',subject_id)
                count=int(request.form.get('count','10'))
                mode=request.form.get('mode','練習')
                if not 1<=count<=100 or mode not in ('練習','模擬考','錯題複習'):
                    raise ValueError('請選擇有效模式與 1–100 題。')
                clause=' AND q.id IN (SELECT question_id FROM wrong_answers WHERE user_id=? AND status=\'待複習\')' if mode=='錯題複習' else ''
                params=(subject_id,g.user['id'],count) if clause else (subject_id,count)
                questions=db().execute('SELECT q.* FROM questions q JOIN chapters c ON c.id=q.chapter_id WHERE c.subject_id=?'+clause+' ORDER BY random() LIMIT ?',params).fetchall()
                if not questions:
                    raise ValueError('此科目尚無可用題目，請先建立題庫或累積錯題。')
                sid=insert('quiz_sessions',dict(user_id=g.user['id'],subject_id=subject_id,mode=mode,total_count=len(questions)))
                for q in questions:
                    snapshot=dict(q)
                    snapshot['options']=[dict(r) for r in db().execute('SELECT * FROM question_options WHERE question_id=? ORDER BY order_no',(q['id'],))]
                    insert('quiz_answers',dict(session_id=sid,question_id=q['id'],question_snapshot=json.dumps(snapshot,ensure_ascii=False)))
                db().commit()
                return redirect(url_for('quiz_take',sid=sid))
            except ValueError as exc:
                db().rollback()
                error=str(exc)
        sessions=db().execute('SELECT * FROM quiz_sessions WHERE user_id=? ORDER BY id DESC',(g.user['id'],)).fetchall()
        return render_template('quiz_start.html',title='模擬考試',subjects=options(db(),'subjects',g.user['id']),sessions=sessions,error=error)

    @app.route('/quiz/<int:sid>',methods=['GET','POST'])
    @login_required
    def quiz_take(sid):
        exam=owned('quiz_sessions',sid)
        if request.method=='POST':
            db().execute('BEGIN IMMEDIATE')
            exam=owned('quiz_sessions',sid)
            if exam['finished_at']:
                db().rollback()
                return redirect(url_for('quiz_take',sid=sid))
            correct=0
            for row in db().execute('SELECT * FROM quiz_answers WHERE session_id=?',(sid,)).fetchall():
                q=json.loads(row['question_snapshot'])
                key='answer_'+str(row['id'])
                answer=','.join(sorted(set(request.form.getlist(key)))) if q['q_type']=='多選' else request.form.get(key,'').strip()
                if len(answer)>50:
                    db().rollback()
                    abort(400,'答案不可超過 50 字。')
                ok=answer==q['answer_key']
                correct+=int(ok)
                db().execute('UPDATE quiz_answers SET user_answer=?,is_correct=?,answered_at=CURRENT_TIMESTAMP WHERE id=?',(answer,int(ok),row['id']))
                if not ok:
                    db().execute("INSERT INTO wrong_answers(user_id,question_id,wrong_count,last_wrong_at,status) VALUES (?,?,1,CURRENT_TIMESTAMP,'待複習') ON CONFLICT(user_id,question_id) DO UPDATE SET wrong_count=wrong_count+1,last_wrong_at=CURRENT_TIMESTAMP,status='待複習'",(g.user['id'],q['id']))
                elif exam['mode']=='錯題複習':
                    db().execute("UPDATE wrong_answers SET status='已克服' WHERE user_id=? AND question_id=?",(g.user['id'],q['id']))
            db().execute('UPDATE quiz_sessions SET correct_count=?,finished_at=CURRENT_TIMESTAMP WHERE id=?',(correct,sid))
            refresh_stats(commit=False)
            db().commit()
            return redirect(url_for('quiz_take',sid=sid))
        items=[]
        for row in db().execute('SELECT * FROM quiz_answers WHERE session_id=? ORDER BY id',(sid,)):
            q=json.loads(row['question_snapshot'])
            if not exam['finished_at']:
                q.pop('answer_key',None)
                q.pop('explanation',None)
            items.append((row,q))
        return render_template('quiz_take.html',title='考試結果' if exam['finished_at'] else '作答中',exam=exam,items=items)

    @app.get('/results')
    @login_required
    def results():
        wrong=db().execute('SELECT w.*,q.content,q.explanation FROM wrong_answers w JOIN questions q ON q.id=w.question_id WHERE w.user_id=? ORDER BY w.last_wrong_at DESC',(g.user['id'],)).fetchall()
        sessions=db().execute('SELECT * FROM quiz_sessions WHERE user_id=? AND finished_at IS NOT NULL ORDER BY id DESC',(g.user['id'],)).fetchall()
        return render_template('results.html',title='成績與錯題簿',wrong=wrong,sessions=sessions)

    @app.get('/analysis')
    @login_required
    def analysis():
        rows=db().execute('SELECT s.subject_name,c.chapter_name,count(*) n,sum(a.is_correct) correct FROM quiz_answers a JOIN quiz_sessions qs ON qs.id=a.session_id JOIN questions q ON q.id=a.question_id JOIN chapters c ON c.id=q.chapter_id JOIN subjects s ON s.id=c.subject_id WHERE qs.user_id=? AND qs.finished_at IS NOT NULL GROUP BY c.id ORDER BY CAST(sum(a.is_correct) AS REAL)/count(*)',(g.user['id'],)).fetchall()
        return render_template('analysis.html',title='章節弱項分析',rows=rows)

    @app.get('/materials')
    @login_required
    def materials():
        return render_template('materials.html',title='教材與 RAG（第二階段）')

    @app.get('/imports/template')
    @login_required
    def import_template():
        output=io.StringIO()
        writer=csv.writer(output)
        writer.writerow(['chapter_name','q_type','content','answer_key','explanation','difficulty','option_A','option_B','option_C','option_D'])
        writer.writerow(['範例章節','單選','下列何者是關聯式資料庫？','A','SQLite 是關聯式資料庫。',1,'SQLite','HTML','',''])
        return Response('\ufeff'+output.getvalue(),mimetype='text/csv',headers={'Content-Disposition':'attachment; filename=questions-template.csv'})

    @app.route('/imports',methods=['GET','POST'])
    @login_required
    def imports():
        error=None
        if request.method=='POST':
            try:
                subject_id=int(request.form.get('subject_id','0'))
                owned('subjects',subject_id)
                upload=request.files.get('file')
                if not upload or not upload.filename:
                    raise ValueError('請選擇 CSV 或 XLSX 檔案。')
                suffix=Path(upload.filename).suffix.lower()
                raw=upload.read()
                if suffix=='.csv':
                    rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
                elif suffix=='.xlsx':
                    from openpyxl import load_workbook
                    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                        if sum(z.file_size for z in archive.infolist())>30*1024*1024:
                            raise ValueError('Excel 解壓後過大，請拆分檔案。')
                    book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
                    try:
                        iterator=book.active.iter_rows(values_only=True)
                        keys=next(iterator)
                        rows=[]
                        for line in iterator:
                            if any(v is not None for v in line):
                                rows.append(dict(zip(keys,line)))
                            if len(rows)>500:
                                raise ValueError('每批最多 500 題。')
                    finally:
                        book.close()
                else:
                    raise ValueError('目前只支援 UTF-8 CSV 或 XLSX。')
                if not 1<=len(rows)<=500:
                    raise ValueError('每批需為 1–500 題。')
                clean=[]
                for i,row in enumerate(rows,2):
                    item={str(k):str(v).strip() if v is not None else '' for k,v in row.items()}
                    if not item.get('chapter_name') or not item.get('content') or not item.get('answer_key'):
                        raise ValueError(f'第 {i} 列缺少章節、題目或答案，請依範本填寫。')
                    if len(item['chapter_name'])>120 or len(item['content'])>10000 or len(item['answer_key'])>50 or len(item.get('explanation',''))>10000:
                        raise ValueError(f'第 {i} 列文字過長。')
                    if item.get('q_type') not in ('單選','多選','是非','填空'):
                        raise ValueError(f'第 {i} 列題型無效。')
                    item['difficulty']=int(item.get('difficulty') or '1')
                    if not 1<=item['difficulty']<=5:
                        raise ValueError(f'第 {i} 列難度須為 1–5。')
                    item,_=validate_question(item,item)
                    clean.append(item)
                folder=Path(app.config['DATABASE']).parent/'imports'
                folder.mkdir(exist_ok=True)
                path=folder/(secrets.token_hex(16)+suffix)
                path.write_bytes(raw)
                batch=insert('exam_imports',dict(user_id=g.user['id'],subject_id=subject_id,file_name=Path(upload.filename).name[:255],file_path=str(path),file_type=suffix[1:],status='待確認',total_rows=len(clean)))
                for item in clean:
                    encoded=json.dumps(item,ensure_ascii=False)
                    insert('import_items',dict(import_id=batch,raw_content=encoded,parsed_json=encoded))
                db().commit()
                return redirect(url_for('import_review',batch=batch))
            except (ValueError,UnicodeError,zipfile.BadZipFile,StopIteration,KeyError,TypeError,OverflowError,ParseError) as exc:
                db().rollback()
                error=str(exc) if isinstance(exc,ValueError) and not isinstance(exc,UnicodeError) else '無法讀取檔案，請使用正確的 UTF-8 CSV / XLSX 範本。'
        batches=db().execute('SELECT * FROM exam_imports WHERE user_id=? ORDER BY id DESC',(g.user['id'],)).fetchall()
        return render_template('imports.html',title='考試資料匯入',subjects=options(db(),'subjects',g.user['id']),batches=batches,error=error)

    @app.route('/imports/<int:batch>',methods=['GET','POST'])
    @login_required
    def import_review(batch):
        record=owned('exam_imports',batch)
        if request.method=='POST':
            db().execute('BEGIN IMMEDIATE')
            record=owned('exam_imports',batch)
            if record['status']=='待確認':
                owned('subjects',record['subject_id'])
                for row in db().execute('SELECT * FROM import_items WHERE import_id=?',(batch,)).fetchall():
                    item=json.loads(row['parsed_json'])
                    chapter=db().execute('SELECT id FROM chapters WHERE subject_id=? AND chapter_name=?',(record['subject_id'],item['chapter_name'])).fetchone()
                    cid=chapter['id'] if chapter else insert('chapters',dict(subject_id=record['subject_id'],chapter_name=item['chapter_name']))
                    data={k:item.get(k,'') for k in ('q_type','content','answer_key','explanation','difficulty')}
                    data,opts=validate_question(data,item)
                    qid=insert('questions',dict(chapter_id=cid,source='import',**data))
                    for label,text in opts:
                        insert('question_options',dict(question_id=qid,option_label=label,option_text=text,order_no=ord(label)))
                    db().execute('UPDATE import_items SET chapter_id=?,question_id=?,is_confirmed=1 WHERE id=?',(cid,qid,row['id']))
                db().execute("UPDATE exam_imports SET status='已匯入',imported_count=total_rows WHERE id=?",(batch,))
            db().commit()
            flash('已確認並匯入題庫。','success')
            return redirect(url_for('imports'))
        items=[json.loads(r['parsed_json']) for r in db().execute('SELECT * FROM import_items WHERE import_id=?',(batch,))]
        return render_template('import_review.html',title='確認匯入題目',batch=record,items=items)
