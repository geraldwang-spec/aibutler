from __future__ import annotations
from pathlib import Path
import base64
import zipfile
import re
from .question_ranges import reading_ranges


def pdf_reading_anchors(page):
    """Locate native-text instructions in reading order, including wrapped lines."""
    anchors=[]
    for block in page.get_text('dict')['blocks']:
        lines=block.get('lines',[])
        text='\n'.join(''.join(span['text'] for span in line['spans']) for line in lines)
        for anchor in reading_ranges(text):
            line_index=text[:anchor['start']].count('\n')
            bbox=lines[line_index]['bbox']
            column=int((bbox[0]+bbox[2])/2>=page.rect.width/2)
            anchors.append(dict(anchor, order=(column,bbox[1]), bbox=list(bbox)))
    return sorted(anchors,key=lambda a:a['order'])


def pdf_question_anchors(page):
    from .question_importer import _question_start
    anchors=[]
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            text=''.join(span['text'] for span in line['spans'])
            number=_question_start(text)
            if number is None: continue
            bbox=line['bbox']
            anchors.append(dict(number=number,text=text,order=(int((bbox[0]+bbox[2])/2>=page.rect.width/2),bbox[1]),bbox=list(bbox)))
    # Formula denominators such as "1)" are not question labels merely
    # because text extraction puts them on a new line. Require alignment
    # with labels accompanied by prose; if none exist, retain for review.
    strong=[a for a in anchors if re.search(r'[\u4e00-\u9fff]{2}|[A-Za-z]{3}',a['text'])]
    if len(strong)>=2:
        anchors=[a for a in anchors if a in strong or any(abs(a['bbox'][0]-s['bbox'][0])<=
                 (s['bbox'][3]-s['bbox'][1])*.5 for s in strong)]
    return sorted(anchors,key=lambda a:a['order'])


def image_scope(anchors, question_anchors, bbox, page_width, carried_range=None):
    from .layout_regions import visual_region
    return visual_region(anchors,question_anchors,bbox,page_width,carried_range)['range']


def image_text_kind(text):
    markers=re.findall(r'(?m)^\s*[（(]?\s*([A-DＡ-Ｄ])(?=[）).．、:：\s]|[\u4e00-\u9fff])',text)
    return 'choices' if len(set(markers))>=2 else 'material'


def image_text_fields(text):
    """Separate OCR choices from material; keep raw source alongside fields."""
    marker=re.compile(r'(?m)^\s*[（(]?\s*([A-DＡ-Ｄ])(?=[）).．、:：\s]|[\u4e00-\u9fff])')
    hits=list(marker.finditer(text))
    labels=[h.group(1).translate(str.maketrans('ＡＢＣＤ','ABCD')) for h in hits]
    if len(hits)<2:
        return [dict(field='material',text=text)]
    if labels!=list('ABCD')[:len(labels)]:
        return [dict(field='unresolved',text=text)]
    fields=[]
    if text[:hits[0].start()].strip(): fields.append(dict(field='material',text=text[:hits[0].start()].strip()))
    for i,hit in enumerate(hits):
        fields.append(dict(field=labels[i],text=text[hit.start():hits[i+1].start() if i+1<len(hits) else len(text)].strip()))
    return fields


def exam_sections(sections):
    """Filter administrative covers for exams only, preserving knowledge sources."""
    result=[]
    for section in sections:
        text=section.get('exam_text',section.get('text',''))
        cues=sum(bool(re.search(pattern,text,re.I)) for pattern in
                 ('注意事項|INSTRUCTIONS TO CANDIDATES','准考證|admission ticket','考試開始|作答方式|應試號碼'))
        choice_lines=re.findall(r'(?m)^\s*(?:[（(][A-DＡ-Ｄ][）)]|[A-DＡ-Ｄ][.．、)])',text)
        asks_question=bool(re.search(r'何者|下列|哪[一個項]|[?？]',text))
        if cues>=2 and len(choice_lines)<2 and not asks_question: continue
        result.append(dict(section,text=text))
    return result


def needs_ocr(text):
    if len(text.strip()) < 40:
        return True
    broken = sum(ch == '\ufffd' or '\ue000' <= ch <= '\uf8ff' for ch in text)
    return broken / max(1, len(text)) > .15 or len(re.findall(r'\(cid:\d+\)', text)) >= 3

def parse_file(path:Path, include_pdf_images=False):
    ext=path.suffix.lower()
    if ext=='.pdf':
        try: import fitz
        except ImportError as exc: raise RuntimeError('解析 PDF 需要 PyMuPDF：pip install pymupdf') from exc
        doc=fitz.open(path)
        sections=[]
        try:
            scanned=0
            carried_range=None
            for i,page in enumerate(doc,1):
                image_context=[]
                anchors=pdf_reading_anchors(page)
                question_anchors=pdf_question_anchors(page)
                text=page.get_text('text').strip()
                exam_text=None
                if not needs_ocr(text):
                    from .question_importer import _question_start
                    accepted={re.sub(r'\s+','',q['text']) for q in question_anchors}
                    lines=[]
                    for line in text.splitlines():
                        if _question_start(line) is not None and re.sub(r'\s+','',line) not in accepted and lines:
                            lines[-1]+=' '+line
                            # Joining "(" + "1)" must not create a fresh
                            # false label "( 1)". Fold until a non-label line.
                            while len(lines)>1 and _question_start(lines[-1]) is not None and re.sub(r'\s+','',lines[-1]) not in accepted:
                                tail=lines.pop(); lines[-1]+=' '+tail
                        else: lines.append(line)
                    exam_text='\n'.join(lines)
                if needs_ocr(text):
                    scanned+=1
                    from .exam_modules import cpu
                    pixels=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
                    text=cpu('ocr',image=base64.b64encode(pixels.tobytes('png')).decode('ascii'))
                elif include_pdf_images and page.get_image_info():
                    # Keep image OCR separate from numbered question text and
                    # attach it to its explicitly labelled reading group.
                    from .exam_modules import cpu
                    infos=sorted(page.get_image_info(),key=lambda x:(int((x['bbox'][0]+x['bbox'][2])/2>=page.rect.width/2),x['bbox'][1]))
                    for info in infos:
                        rect=fitz.Rect(info['bbox']) & page.rect
                        if rect.width<35 or rect.height<25: continue
                        scanned+=1
                        pixels=page.get_pixmap(matrix=fitz.Matrix(2,2),clip=rect,alpha=False)
                        image_text=cpu('ocr',image=base64.b64encode(pixels.tobytes('png')).decode('ascii'))
                        # OCR list numbers/time values must never become exam
                        # question numbers. Keep image text in separate metadata.
                        from .layout_regions import visual_region
                        scope=visual_region(anchors,question_anchors,rect,page.rect.width,carried_range)
                        limits=scope['range']
                        image_anchors=reading_ranges(image_text)
                        if image_anchors:
                            own=image_anchors[-1]
                            limits=(own['question_from'],own['question_to'])
                        if image_text.strip():
                            fields=image_text_fields(image_text)
                            material='\n'.join(f['text'] for f in fields if f['field']=='material')
                            if any(f['field']=='unresolved' for f in fields):
                                scope=dict(scope,status='unresolved',reason='ambiguous_ocr_fields')
                                limits=None
                            image_context.append(dict(scope,range=limits,text=material,source_text=image_text,fields=fields,
                                                      bbox=list(rect),kind='material' if material else 'choices'))
                if anchors:
                    carried_range=(anchors[-1]['question_from'],anchors[-1]['question_to'])
                if text: sections.append({'locator':f'page:{i}','title':f'第 {i} 頁','text':text,'exam_text':exam_text or text,'image_context':image_context,
                                          'reading_groups':anchors,'has_images':bool(page.get_image_info()),
                                          'text_layer_unreliable':sum(ch=='\ufffd' or '\ue000'<=ch<='\uf8ff' for ch in text)>=3})
        finally:
            doc.close()
        if not sections: raise RuntimeError('PDF 文字與 CPU OCR 均未取得文字，請檢查掃描品質。')
        return sections
    if ext=='.docx':
        try: from docx import Document
        except ImportError as exc: raise RuntimeError('解析 Word 需要 python-docx。') from exc
        doc=Document(path); sections=[]; buf=[]; sec=1
        from docx.table import Table
        table_no = 0
        for p in doc.iter_inner_content():
            if isinstance(p, Table):
                if buf:
                    sections.append({'locator':f'section:{sec}','title':buf[0][:80],'text':'\n'.join(buf)}); sec+=1; buf=[]
                table_no += 1
                rows=['\t'.join(cell.text for cell in row.cells) for row in p.rows]
                sections.append(dict(locator=f'table:{table_no}',title=f'表格 {table_no}',text='\n'.join(rows)))
                continue
            t=p.text.strip()
            if not t: continue
            if p.style and str(p.style.name).lower().startswith('heading') and buf:
                sections.append({'locator':f'section:{sec}','title':buf[0][:80],'text':'\n'.join(buf)}); sec+=1; buf=[]
            buf.append(t)
        if buf: sections.append({'locator':f'section:{sec}','title':buf[0][:80],'text':'\n'.join(buf)})
        return sections
    if ext in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        out=[]
        for ws in wb.worksheets:
            rows=[]
            for row in ws.iter_rows(values_only=True):
                vals=['' if v is None else str(v) for v in row]
                if any(vals): rows.append('\t'.join(vals))
            if rows: out.append({'locator':f'sheet:{ws.title}','title':ws.title,'text':'\n'.join(rows)})
        wb.close(); return out
    if ext in ('.txt','.md','.csv'):
        raw=path.read_bytes()
        text=None
        for encoding in ('utf-8-sig','utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'cp950'):
            try:
                text=raw.decode(encoding); break
            except UnicodeError:
                pass
        if text is None: raise RuntimeError('無法辨識文字編碼，請以 UTF-8 儲存後重試。')
        return [{'locator':'text:1','title':path.name,'text':text}]
    raise RuntimeError('目前支援 PDF、DOCX、XLSX、TXT、MD、CSV。')

def chunk_sections(sections, size=1400, overlap=180):
    chunks=[]
    for sec in sections:
        text=' '.join(sec['text'].split())
        start=0
        while start < len(text):
            end=min(len(text),start+size)
            part=text[start:end].strip()
            if part: chunks.append({'text':part,'locator':sec['locator'],'title':sec['title']})
            if end>=len(text): break
            start=max(start+1,end-overlap)
    return chunks
